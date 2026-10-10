import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_structured_abstract(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup with clean 1-inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    BLACK = RGBColor(0, 0, 0)
    DARK_BLUE = RGBColor(20, 40, 80)

    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = BLACK

    # 1. Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.paragraph_format.line_spacing = 1.15

    r_title = p_title.add_run("EDUMATRIX: INSTITUTIONAL ACADEMIC AND PLACEMENT INTELLIGENCE SYSTEM")
    r_title.font.name = 'Times New Roman'
    r_title.bold = True
    r_title.font.size = Pt(13.5)
    r_title.font.color.rgb = DARK_BLUE

    # 2. Main Heading: Structured Abstract
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.paragraph_format.space_before = Pt(2)
    p_head.paragraph_format.space_after = Pt(16)

    r_head = p_head.add_run("Structured Abstract")
    r_head.font.name = 'Times New Roman'
    r_head.bold = True
    r_head.underline = True
    r_head.font.size = Pt(12.5)
    r_head.font.color.rgb = BLACK

    # Helper function to add a structured abstract section
    def add_section(label, text, space_after=10):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15

        # Bold Label
        r_lbl = p.add_run(f"•  {label}: ")
        r_lbl.font.name = 'Times New Roman'
        r_lbl.bold = True
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.color.rgb = BLACK

        # Content Text
        r_txt = p.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10.5)
        r_txt.font.color.rgb = BLACK
        return p

    # 1. Purpose
    add_section(
        "Purpose",
        "Higher educational institutions frequently struggle with fragmented communication across informal platforms like WhatsApp, email threads, and notice boards, causing administrative delays, confusion, and missed deadlines. Furthermore, manual placement tracking fails to scale efficiently with large student datasets. This research develops EDUMATRIX, a centralized, AI-powered academic and placement intelligence platform designed to streamline real-time institutional communication and automate placement workflows."
    )

    # 2. Design/methodology/approach
    add_section(
        "Design/methodology/approach",
        "EDUMATRIX is implemented using a reactive full-stack architecture comprising React, FastAPI, PostgreSQL, Redis, and WebSockets, enforced with role-based access control (RBAC) across student, faculty, and administrative portals. The system embeds machine learning and NLP pipelines, utilizing Named Entity Recognition (NER) to parse event names, dates, and deadlines from uploaded PDF circulars, alongside models for intent classification, sentiment analysis, placement readiness prediction, and skill-to-company matching."
    )

    # 3. Findings
    add_section(
        "Findings",
        "Empirical deployment demonstrates that centralizing communication significantly reduces information bottlenecks and manual overhead. Automated NER-driven circular parsing eliminates human error in scheduling notices, while the predictive placement engine provides objective readiness scoring and accurate skill matching to enhance student career outcomes."
    )

    # 4. Research limitations/implications
    add_section(
        "Research limitations/implications",
        "The current system is evaluated within a single institutional environment using standard circular document formats and structured academic datasets. Future extensions will investigate cross-institutional federated deployments, multi-lingual circular extraction, and adaptive mock-interview evaluation."
    )

    # 5. Originality/value
    add_section(
        "Originality/value",
        "EDUMATRIX addresses the critical gap in educational management by coupling real-time collaborative communication with document intelligence and predictive placement analytics in a unified, open-source-driven institutional intelligence ecosystem.",
        space_after=14
    )

    # Keywords
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(4)
    p_kw.paragraph_format.space_after = Pt(0)
    p_kw.paragraph_format.line_spacing = 1.15

    r_kw_lbl = p_kw.add_run("Keywords: ")
    r_kw_lbl.font.name = 'Times New Roman'
    r_kw_lbl.bold = True
    r_kw_lbl.font.size = Pt(10)
    r_kw_lbl.font.color.rgb = BLACK

    r_kw_val = p_kw.add_run(
        "Institutional Intelligence, Academic Management, Placement Readiness, Named Entity Recognition (NER), Machine Learning, FastAPI, WebSockets."
    )
    r_kw_val.font.name = 'Times New Roman'
    r_kw_val.font.size = Pt(10)
    r_kw_val.font.color.rgb = BLACK

    doc.save(output_path)
    print(f"Structured abstract created at: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence\EDUMATRIX_Structured_Abstract.docx"
    generate_structured_abstract(out_file)
