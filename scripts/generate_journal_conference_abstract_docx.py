import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_doc(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Color Palette - Professional Academic Journal Tone
    NAVY = RGBColor(20, 40, 80)       # Dark Academic Navy
    CHARCOAL = RGBColor(35, 35, 35)   # Body Text
    MUTED = RGBColor(90, 95, 105)     # Conference / Affiliation Metadata

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = CHARCOAL

    # 1. CONFERENCE / JOURNAL HEADER (Top-Right or Centered subtle tag)
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(10)
    r_hdr = p_header.add_run("INTERNATIONAL CONFERENCE / JOURNAL SUBMISSION")
    r_hdr.font.name = 'Times New Roman'
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = MUTED

    # 2. TITLE
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.paragraph_format.line_spacing = 1.18

    r_title = p_title.add_run(
        "EDUMATRIX: AN AI-POWERED INSTITUTIONAL ACADEMIC AND\n"
        "PLACEMENT INTELLIGENCE SYSTEM"
    )
    r_title.font.name = 'Times New Roman'
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = NAVY

    # 3. AUTHOR / AFFILIATION BLOCK
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(0)
    p_auth.paragraph_format.space_after = Pt(14)
    p_auth.paragraph_format.line_spacing = 1.15

    r_auth = p_auth.add_run("Author / Research Team\n")
    r_auth.font.name = 'Times New Roman'
    r_auth.font.size = Pt(10)
    r_auth.font.bold = True
    r_auth.font.color.rgb = CHARCOAL

    r_affil = p_auth.add_run("Department of Computer Science and Engineering\nInstitution / University Academic Research")
    r_affil.font.name = 'Times New Roman'
    r_affil.font.size = Pt(9.5)
    r_affil.font.italic = True
    r_affil.font.color.rgb = MUTED

    # 4. ABSTRACT HEADING
    p_abs_hd = doc.add_paragraph()
    p_abs_hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_hd.paragraph_format.space_before = Pt(4)
    p_abs_hd.paragraph_format.space_after = Pt(8)

    r_abs_hd = p_abs_hd.add_run("ABSTRACT")
    r_abs_hd.font.name = 'Times New Roman'
    r_abs_hd.bold = True
    r_abs_hd.font.size = Pt(11)
    r_abs_hd.font.color.rgb = NAVY

    # Helper for justified body paragraph
    def add_body_paragraph(text, space_after=7):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.color.rgb = CHARCOAL
        return p

    # PARAGRAPH 1: PROBLEM STATEMENT & CENTRALIZED PLATFORM PROPOSAL
    add_body_paragraph(
        "EDUMATRIX is an AI-powered Institutional Academic and Placement Intelligence System designed to simplify and improve communication, academic management, and placement activities within colleges and educational institutions. In many institutions, important information related to academics and placements is still shared through multiple platforms such as WhatsApp groups, emails, and notice boards. This often creates communication gaps, confusion, and missed deadlines among students. Managing placement activities manually also becomes difficult when handling large amounts of student data. To address these problems, EDUMATRIX provides a centralized and intelligent platform that combines real-time communication, automated notifications, and AI-based analysis into a single system."
    )

    # PARAGRAPH 2: TECHNICAL ARCHITECTURE & ROLE-BASED ACCESS
    add_body_paragraph(
        "The system is developed using modern technologies such as React, FastAPI, PostgreSQL, Redis, and WebSocket communication to ensure smooth performance and real-time interaction. It provides separate dashboards for students, lecturers, and administrators with secure role-based access. The platform supports features such as real-time chat, circular management, push notifications, academic reminders, and placement updates, helping institutions manage academic activities more efficiently."
    )

    # PARAGRAPH 3: AI/NLP MODULES & DATA EXTRACTION
    add_body_paragraph(
        "Artificial Intelligence and Natural Language Processing techniques are integrated into the system to automate several important tasks. The AI modules perform intent classification, sentiment analysis, placement readiness prediction, skill matching, and Named Entity Recognition (NER). The NER model automatically extracts important details such as event names, dates, deadlines, and announcements from uploaded PDF circulars. The system also analyzes student academic performance, technical skills, and certifications to generate placement readiness insights and company matching recommendations."
    )

    # PARAGRAPH 4: IMPACT & CONCLUSION
    add_body_paragraph(
        "By combining AI-driven automation with real-time communication technologies, EDUMATRIX reduces manual effort, improves information accessibility, minimizes communication delays, and enhances overall institutional management. The project demonstrates how Artificial Intelligence and Machine Learning can be effectively applied in the education sector to create smarter, more organized, and technology-driven academic environments.",
        space_after=14
    )

    # 5. KEYWORDS SECTION
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(4)
    p_kw.paragraph_format.space_after = Pt(0)
    p_kw.paragraph_format.line_spacing = 1.15

    r_kw_label = p_kw.add_run("Keywords— ")
    r_kw_label.font.name = 'Times New Roman'
    r_kw_label.bold = True
    r_kw_label.font.size = Pt(9.5)
    r_kw_label.font.color.rgb = NAVY

    r_kw_val = p_kw.add_run(
        "Academic Management System, Placement Intelligence, Named Entity Recognition (NER), Machine Learning, FastAPI, React, WebSockets, Skill Matching, Sentiment Analysis, Natural Language Processing."
    )
    r_kw_val.font.name = 'Times New Roman'
    r_kw_val.font.size = Pt(9.5)
    r_kw_val.font.color.rgb = CHARCOAL

    doc.save(output_path)
    print(f"Generated successfully at: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence\EDUMATRIX_Journal_Conference_Abstract.docx"
    generate_doc(out_file)
