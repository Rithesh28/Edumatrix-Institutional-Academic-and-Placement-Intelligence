import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_human_abstract(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup with clean 1-inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    BLACK = RGBColor(0, 0, 0)

    # Base Style: Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = BLACK

    # 1. Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.line_spacing = 1.2

    r_title = p_title.add_run("EDUMATRIX: INSTITUTIONAL ACADEMIC AND PLACEMENT INTELLIGENCE SYSTEM")
    r_title.font.name = 'Times New Roman'
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = BLACK

    # 2. Heading: ABSTRACT
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.paragraph_format.space_before = Pt(6)
    p_head.paragraph_format.space_after = Pt(20)

    r_head = p_head.add_run("ABSTRACT")
    r_head.font.name = 'Times New Roman'
    r_head.bold = True
    r_head.font.size = Pt(13)
    r_head.font.color.rgb = BLACK

    # 3. Single Unified Human Paragraph
    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.space_before = Pt(0)
    p_body.paragraph_format.space_after = Pt(0)
    p_body.paragraph_format.line_spacing = 1.4  # Clean academic line spacing

    body_text = (
        "In many colleges and institutions, academic circulars and placement notices are scattered across WhatsApp "
        "groups, email threads, and physical notice boards, leading to communication gaps, confusion, and missed "
        "deadlines among students. Managing placement data manually also becomes difficult and time-consuming as "
        "student numbers grow. To address these problems, EDUMATRIX was developed as a centralized web platform "
        "combining real-time communication, automated notifications, and AI-driven analytics into a single system. "
        "The platform is built using a modern stack comprising React for the frontend, FastAPI for the backend, "
        "PostgreSQL for data storage, and Redis with WebSockets for real-time messaging, with separate role-based "
        "access for students, lecturers, and administrators. Natural Language Processing (NLP) techniques, including "
        "Named Entity Recognition (NER), are implemented to automatically extract event titles, dates, and deadlines "
        "directly from uploaded PDF circulars, alongside machine learning modules for student skill matching, "
        "sentiment analysis, and placement readiness evaluation. Practical testing of the platform demonstrates that "
        "centralizing communication and automating circular data extraction significantly minimizes notification delays, "
        "cuts down administrative workload, and provides clear, data-driven placement insights for students. While the "
        "current implementation is tested within an institutional setup using standard circular templates, future work "
        "can expand the system to support multi-college networks and multilingual document parsing. Overall, EDUMATRIX "
        "offers practical value by bringing daily academic communication and placement intelligence together into a single, "
        "organized, and reliable platform."
    )

    r_body = p_body.add_run(body_text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = BLACK

    doc.save(output_path)
    print(f"Human abstract generated successfully at: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence\EDUMATRIX_Human_Abstract.docx"
    generate_human_abstract(out_file)
