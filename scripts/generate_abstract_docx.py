import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_clean_abstract_docx(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # Professional academic color scheme
    NAVY = RGBColor(15, 34, 64)       # Deep Institutional Navy for Title
    DARK_TEXT = RGBColor(26, 26, 26)  # Near-black for maximum readability
    ACCENT = RGBColor(30, 60, 100)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = DARK_TEXT

    # 1. TITLE
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.paragraph_format.line_spacing = 1.15

    r_title = p_title.add_run("EDUMATRIX: INSTITUTIONAL ACADEMIC & PLACEMENT\nINTELLIGENCE SYSTEM")
    r_title.bold = True
    r_title.font.size = Pt(13.5)
    r_title.font.color.rgb = NAVY

    # 2. ABSTRACT HEADER
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(4)
    p_abs.paragraph_format.space_after = Pt(8)

    r_abs = p_abs.add_run("ABSTRACT")
    r_abs.bold = True
    r_abs.font.size = Pt(11)
    r_abs.font.color.rgb = NAVY

    # Helper for body paragraphs
    def add_p(text, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15

        r = p.add_run(text)
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_TEXT
        return p

    # PARAGRAPH 1: PROBLEM STATEMENT & CONTEXT
    add_p(
        "Higher educational institutions face persistent operational friction due to disparate communication channels, manual processing of administrative notices, scheduling conflicts in faculty examination invigilation, and an absence of data-driven tracking for student career readiness. Official academic circulars and critical deadlines are frequently misplaced within unstructured email threads and messaging groups, leading to missed student submissions. In parallel, institutional placement cells struggle with labor-intensive candidate evaluation, lacking automated tools to objectively analyze resumes, match candidate skill profiles against dynamic employer job specifications, or identify actionable domain skill deficiencies."
    )

    # PARAGRAPH 2: PROPOSED EDUMATRIX PLATFORM & ARCHITECTURE
    add_p(
        "To resolve these administrative and placement challenges, this project implements EduMatrix, an end-to-end institutional intelligence platform. Developed on a decoupled, asynchronous micro-architecture, the backend is built with FastAPI and PostgreSQL for robust data management, paired with a high-dimensional Qdrant vector database for semantic retrieval, while the presentation tier features a responsive React web interface along with cross-platform mobile deployment via Capacitor. The system implements comprehensive Role-Based Access Control (RBAC) customized for Students, Lecturers, Heads of Departments (HOD), Placement Officers, and Administrators, ensuring structured data segregation and secure institutional workflows."
    )

    # PARAGRAPH 3: CORE IMPLEMENTED SUBSYSTEMS & TECHNICAL METHODOLOGY
    add_p(
        "The implemented system unifies several core modules: (1) Smart Circular & Knowledge Ingestion: Multi-format circulars (PDFs, scanned images) are processed using OCR and Named Entity Recognition (NER) to extract critical event dates and automatically populate personal reminder schedules, with textual content embedded via Sentence-Transformers (all-MiniLM-L6-v2) into Qdrant for natural language search; (2) Automated Faculty Examination Duty Scheduler: A constraint-satisfaction scheduling engine parses departmental exam rosters, detects venue, date, and session conflicts, and ensures balanced invigilation load distribution across faculty members; (3) Placement Management & AI Resume Analytics: An intelligent recruitment portal supporting recruiter job postings with dynamic eligibility filters, candidate application tracking, and an automated resume analyzer. The analyzer performs multi-category technical and soft skill extraction, calculates a multi-factor ATS score across skills, education, experience, and projects, performs cosine-similarity matching against target job descriptions, and provides structured skill-gap feedback; and (4) Context-Aware Conversational Assistant: An NLP-driven helpdesk utilizing intent classification and vector retrieval to resolve student and faculty queries regarding institutional policies."
    )

    # PARAGRAPH 4: SYSTEM INTEGRATION, RESULTS & FUTURE ROADMAP
    add_p(
        "System operations are supported by asynchronous background tasks (Celery/Redis) and real-time WebSocket notifications for broadcast notices and duty updates. System testing verifies reliable OCR accuracy, zero-conflict exam duty allocations, sub-second vector search response times, and consistent candidate-job relevance scoring. Ongoing and future milestones include expanding the system with predictive placement probability modeling using fine-tuned language models, automated AI mock interview evaluations, and cross-departmental academic performance analytics to further enrich institutional decision-making.",
        space_after=10
    )

    # KEYWORDS
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(4)
    p_kw.paragraph_format.space_after = Pt(0)
    p_kw.paragraph_format.line_spacing = 1.15

    r_kw_lbl = p_kw.add_run("Keywords: ")
    r_kw_lbl.bold = True
    r_kw_lbl.font.size = Pt(9.5)
    r_kw_lbl.font.color.rgb = NAVY

    r_kw_val = p_kw.add_run(
        "Institutional Intelligence, Natural Language Processing, Resume Parsing, Applicant Tracking System (ATS), Semantic Vector Search, Examination Scheduling, FastAPI, Qdrant, Role-Based Access Control."
    )
    r_kw_val.font.size = Pt(9.5)
    r_kw_val.font.color.rgb = DARK_TEXT

    doc.save(output_path)
    print(f"Clean abstract saved to {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence"
    out_file = os.path.join(out_dir, "EduMatrix_Project_Abstract.docx")
    create_clean_abstract_docx(out_file)
