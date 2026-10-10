import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('EDUMATRIX_Research_Paper.docx')
for i in range(30, 45):
    if i < len(doc.paragraphs):
        print(f"P{i}: '{doc.paragraphs[i].text}'")
