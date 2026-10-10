import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_clean_structured_abstract(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup with 1-inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    BLACK = RGBColor(0, 0, 0)

    # Base Normal Style: Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = BLACK

    # 1. Title Only (Centered, Bold, 14pt)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.line_spacing = 1.2

    r_title = p_title.add_run("EDUMATRIX: INSTITUTIONAL ACADEMIC AND PLACEMENT INTELLIGENCE SYSTEM")
    r_title.font.name = 'Times New Roman'
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = BLACK

    # 2. Heading: ABSTRACT (Centered, Bold, 13pt)
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(4)
    p_abs.paragraph_format.space_after = Pt(18)

    r_abs = p_abs.add_run("STRUCTURED ABSTRACT")
    r_abs.font.name = 'Times New Roman'
    r_abs.bold = True
    r_abs.font.size = Pt(13)
    r_abs.font.color.rgb = BLACK

    # Helper function for each structured section
    def add_section(heading, text, space_after=12):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25

        r_h = p.add_run(heading + ": ")
        r_h.font.name = 'Times New Roman'
        r_h.bold = True
        r_h.font.size = Pt(11)
        r_h.font.color.rgb = BLACK

        r_t = p.add_run(text)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = BLACK
        return p

    # 1. Purpose
    add_section(
        "Purpose",
        "In many colleges and institutions, academic circulars and placement notices are scattered across WhatsApp "
        "groups, email threads, and physical notice boards, leading to communication gaps, confusion, and missed deadlines "
        "among students. Managing placement data manually also becomes difficult and time-consuming as student numbers grow. "
        "To address these problems, EDUMATRIX was developed as a centralized web platform combining real-time communication, "
        "automated notifications, and AI-driven analytics into a single system."
    )

    # 2. Methodology
    add_section(
        "Methodology",
        "The platform is built using a modern stack comprising React for the frontend, FastAPI for the backend, "
        "PostgreSQL for data storage, and Redis with WebSockets for real-time messaging, with separate role-based access "
        "for students, lecturers, and administrators. Natural Language Processing (NLP) techniques, including Named Entity "
        "Recognition (NER), are implemented to automatically extract event titles, dates, and deadlines directly from uploaded "
        "PDF circulars, alongside machine learning modules for student skill matching, sentiment analysis, and placement "
        "readiness evaluation."
    )

    # 3. Findings & Research Implications
    add_section(
        "Findings & Research Implications",
        "Practical testing of the platform demonstrates that centralizing communication and automating circular data "
        "extraction significantly minimizes notification delays, cuts down administrative workload, and provides clear, "
        "data-driven placement insights for students. While the current implementation is tested within an institutional "
        "setup using standard circular templates, future work can expand the system to support multi-college networks and "
        "multilingual document parsing."
    )

    # 4. Originality or Novelty of the Study
    add_section(
        "Originality or Novelty of the Study",
        "Overall, EDUMATRIX offers practical originality and value by bringing daily academic communication, intelligent "
        "document parsing, and placement analytics together into a single, organized, and reliable institutional platform.",
        space_after=0
    )

    doc.save(output_path)
    print(f"Clean structured abstract generated at: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence\EDUMATRIX_Clean_Structured_Abstract.docx"
    generate_clean_structured_abstract(out_file)
