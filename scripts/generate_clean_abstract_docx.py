import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_report_abstract(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup with classic academic report margins (1 inch / 2.54 cm)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Base Normal Style: Classic Times New Roman, Black
    BLACK = RGBColor(0, 0, 0)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = BLACK

    # 1. PROJECT NAME
    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_proj.paragraph_format.space_before = Pt(0)
    p_proj.paragraph_format.space_after = Pt(12)
    p_proj.paragraph_format.line_spacing = 1.2

    r_proj = p_proj.add_run("EDUMATRIX: INSTITUTIONAL ACADEMIC AND PLACEMENT INTELLIGENCE SYSTEM")
    r_proj.font.name = 'Times New Roman'
    r_proj.bold = True
    r_proj.font.size = Pt(14)
    r_proj.font.color.rgb = BLACK

    # 2. ABSTRACT HEADING
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(6)
    p_abs.paragraph_format.space_after = Pt(18)

    r_abs = p_abs.add_run("ABSTRACT")
    r_abs.font.name = 'Times New Roman'
    r_abs.bold = True
    r_abs.font.size = Pt(13)
    r_abs.font.color.rgb = BLACK

    # Helper function for justified body paragraphs
    def add_p(text, space_after=10):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2  # clean academic spacing

        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
        return p

    # PARAGRAPH 1
    add_p(
        "EDUMATRIX is an AI-powered Institutional Academic and Placement Intelligence System designed to "
        "simplify and improve communication, academic management, and placement activities within colleges and "
        "educational institutions. In many institutions, important information related to academics and placements "
        "is still shared through multiple platforms such as WhatsApp groups, emails, and notice boards. This often "
        "creates communication gaps, confusion, and missed deadlines among students. Managing placement activities "
        "manually also becomes difficult when handling large amounts of student data. To address these problems, "
        "EDUMATRIX provides a centralized and intelligent platform that combines real-time communication, "
        "automated notifications, and AI-based analysis into a single system."
    )

    # PARAGRAPH 2
    add_p(
        "The system is developed using modern technologies such as React, FastAPI, PostgreSQL, Redis, and "
        "WebSocket communication to ensure smooth performance and real-time interaction. It provides separate "
        "dashboards for students, lecturers, and administrators with secure role-based access. The platform supports "
        "features such as real-time chat, circular management, push notifications, academic reminders, and "
        "placement updates, helping institutions manage academic activities more efficiently."
    )

    # PARAGRAPH 3
    add_p(
        "Artificial Intelligence and Natural Language Processing techniques are integrated into the system to "
        "automate several important tasks. The AI modules perform intent classification, sentiment analysis, "
        "placement readiness prediction, skill matching, and Named Entity Recognition (NER). The NER model "
        "automatically extracts important details such as event names, dates, deadlines, and announcements from "
        "uploaded PDF circulars. The system also analyzes student academic performance, technical skills, and "
        "certifications to generate placement readiness insights and company matching recommendations."
    )

    # PARAGRAPH 4
    add_p(
        "By combining AI-driven automation with real-time communication technologies, EDUMATRIX reduces manual effort, "
        "improves information accessibility, minimizes communication delays, and enhances overall institutional "
        "management. The project demonstrates how Artificial Intelligence and Machine Learning can be effectively "
        "applied in the education sector to create smarter, more organized, and technology-driven academic environments.",
        space_after=0
    )

    doc.save(output_path)
    print(f"Generated clean project report abstract at: {output_path}")

if __name__ == "__main__":
    out_file = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence\EDUMATRIX_Abstract.docx"
    generate_report_abstract(out_file)
