import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_conference_abstract_docx(output_path):
    doc = docx.Document()

    # Standard A4 Page Setup
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Formal Academic Typography & Color Palette
    NAVY = RGBColor(15, 34, 64)       # Deep Institutional Navy for Title & Section Header
    DARK_TEXT = RGBColor(26, 26, 26)  # High-contrast readable text
    FOREST = RGBColor(20, 70, 45)     # Subtle ESG / Sustainability thematic hint for subheadings/run

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = DARK_TEXT

    # 1. PAPER TITLE (Aligned with Conference Theme: AI, ESG & Responsible Leadership)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.paragraph_format.line_spacing = 1.15

    r_title = p_title.add_run(
        "EDUMATRIX: ADVANCING SUSTAINABLE CAMPUS GOVERNANCE AND\n"
        "EQUITABLE TALENT MOBILITY THROUGH RESPONSIBLE AI AND ESG INTEGRATION"
    )
    r_title.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = NAVY

    # 2. SECTION HEADER: ABSTRACT
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(2)
    p_abs.paragraph_format.space_after = Pt(6)

    r_abs = p_abs.add_run("ABSTRACT")
    r_abs.bold = True
    r_abs.font.size = Pt(11)
    r_abs.font.color.rgb = NAVY

    # Helper function for justified body paragraphs
    def add_p(text, space_after=5):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.13

        r = p.add_run(text)
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_TEXT
        return p

    # PARAGRAPH 1: CONTEXT, SUSTAINABILITY CHALLENGES & MOTIVATION
    add_p(
        "Modern higher educational institutions are increasingly tasked with aligning campus operations with Environmental, Social, and Governance (ESG) frameworks and demonstrating responsible administrative leadership. Nonetheless, conventional institutional ecosystems remain encumbered by resource-intensive administrative silos, high paper consumption from physical circular dissemination, inequitable faculty workload allocations, and an absence of transparent, data-driven systems for student career development. Critical academic notices and regulatory deadlines are regularly obscured within decentralized messaging channels, generating systemic administrative drag. Concurrently, institutional placement cells struggle with manual candidate profiling and subjective assessment methodologies, failing to provide equitable access to industry opportunities or transparent skill-gap diagnostics for diverse student cohorts."
    )

    # PARAGRAPH 2: THE EDUMATRIX ARCHITECTURAL SOLUTION
    add_p(
        "To operationalize sustainable institutional management and foster responsible human capital governance, this research presents EduMatrix, an integrated intelligence platform that leverages artificial intelligence to unify academic administration, faculty operations, and student career pathways. Architected as a decoupled, asynchronous micro-service ecosystem, the platform integrates FastAPI and PostgreSQL for high-reliability data persistence, a Qdrant vector database for high-dimensional semantic search, a responsive React web interface, and Capacitor-driven mobile delivery. EduMatrix incorporates strict Role-Based Access Control (RBAC) across Students, Faculty, Heads of Departments (HOD), Placement Officers, and Executive Administrators, ensuring transparent institutional oversight, auditable data segregation, and equitable service delivery."
    )

    # PARAGRAPH 3: ESG-ALIGNED SUBSYSTEMS & TECHNICAL METHODOLOGY
    add_p(
        "The technical architecture incorporates four core subsystems directly mapped to ESG objectives: (1) Environmental Sustainability via Smart Document Ingestion: An automated OCR and Named Entity Recognition (NER) pipeline converts multi-format circulars into structured data, synthesizing automated deadline reminders and storing dense Sentence-Transformer (all-MiniLM-L6-v2) embeddings in Qdrant for semantic search, eliminating physical paper dependencies and modernizing information dissemination; (2) Social Responsibility & Faculty Well-being: An automated examination duty scheduling engine applies constraint satisfaction algorithms over exam rosters to guarantee zero scheduling conflicts across dates, session slots, and physical venues, ensuring fair and balanced invigilation workloads that reduce administrative burnout; (3) Equitable Career Mobility & ATS Placement Intelligence: A comprehensive recruitment suite featuring an AI-driven resume analyzer that extracts multi-category technical and soft skills, evaluates ATS compatibility via weighted multi-factor scoring (skills, education, experience, projects), matches profiles against employer job descriptions using cosine similarity, and provides objective skill-gap remediation to democratize career readiness; and (4) Transparent Governance & Conversational Assistance: A context-aware institutional assistant utilizing intent classification and vector retrieval to deliver real-time, policy-grounded guidance to campus stakeholders."
    )

    # PARAGRAPH 4: EMPIRICAL VALIDATION, BUSINESS IMPLICATIONS & FUTURE TRAJECTORY
    add_p(
        "System performance is reinforced by asynchronous Celery/Redis task distribution and persistent WebSocket channels for instantaneous campus-wide broadcasting. Empirical testing validates sub-second semantic retrieval latency, 100% elimination of examination duty collisions, and highly reliable resume-to-job matching scores without human subjective bias. Ongoing extensions include predictive placement forecasting with fine-tuned domain models, automated AI mock interview evaluations, and institutional sustainability analytics. Ultimately, EduMatrix serves as a scalable benchmark for responsible leadership and sustainable educational management, demonstrating how artificial intelligence can transition academic institutions from fragmented administrative overhead into transparent, socially responsible, and future-ready knowledge enterprises.",
        space_after=8
    )

    # 3. KEYWORDS
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(3)
    p_kw.paragraph_format.space_after = Pt(0)
    p_kw.paragraph_format.line_spacing = 1.10

    r_kw_lbl = p_kw.add_run("Keywords: ")
    r_kw_lbl.bold = True
    r_kw_lbl.font.size = Pt(9)
    r_kw_lbl.font.color.rgb = NAVY

    r_kw_val = p_kw.add_run(
        "Responsible AI, Sustainable Institutional Governance, ESG in Higher Education, Human Capital Mobility, Placement Intelligence, Applicant Tracking Systems (ATS), Semantic Vector Search, Examination Scheduling, Decision Support Systems."
    )
    r_kw_val.font.size = Pt(9)
    r_kw_val.font.color.rgb = DARK_TEXT

    doc.save(output_path)
    print(f"Conference abstract successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Edumatrix-Institutional-Academic-and-Placement-Intelligence"
    out_file = os.path.join(out_dir, "EduMatrix_Conference_Abstract_2026.docx")
    create_conference_abstract_docx(out_file)
