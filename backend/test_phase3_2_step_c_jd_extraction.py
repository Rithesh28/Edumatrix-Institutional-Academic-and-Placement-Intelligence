"""
Automated Test Suite for Phase 3.2 Step C: General-Purpose JD Extraction Improvement

Comprehensive coverage of the hybrid extraction pipeline:
1. Anonymized Campus Placement Notice (staged comp, multi-percentage, incomplete deadline)
2. Formal Corporate JD (experience, hybrid work mode, conventional salary range, complete deadline)
3. Multi-tier Academic Percentages (10th, 12th, Degree) without bogus CGPA conversion
4. CGPA Requirement Expressions (CGPA >= 7.5, minimum 7.0 CGPA)
5. Table-based Eligibility Extraction
6. Incomplete Date Handling (missing year flagged, no hallucinated year) vs Complete Date
7. Staged Compensation vs Conventional Salary Range (no min-max conflation)
8. Backlog Criteria Parsing (zero backlogs, 1 active arrear)
9. Service Agreement / Bond Extraction
10. Skills Embedded in Natural Language Paragraphs
11. Degree Eligibility and Batch Years
12. Selection Process Rounds Extraction
13. Missing and Ambiguous Fields with Audit Flags
14. PDF Document Ingestion (real PDF generated via reportlab)
15. DOCX Document Ingestion with Tables (real DOCX generated via python-docx)
16. Plain TXT Document Ingestion
17. Empty / Malformed Document Error Handling
18. End-to-end API Parsing via /placement/jobs/parse-jd
19. End-to-end API File Upload Parsing via /placement/jobs/upload-jd
20. DRAFT Job Creation Workflow using Parsed Extracted Values
"""

import os
import sys
import io
import unittest
from datetime import datetime
from unittest.mock import MagicMock, patch

# Configure SQLite for test environment before importing app
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["REDIS_URL"] = "redis://localhost:6379/0"

# Add backend directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.services.jd_parser_service import jd_parser_service
from app.schemas.placement import JdParseResponse, JobPostingCreate
from app.models.job_posting import JobPosting, JobStatus
from app.models.user import User, UserRole


class TestPhase32StepCJdExtraction(unittest.TestCase):
    """Regression test suite for generalized JD extraction."""

    # -------------------------------------------------------------------------
    # TEST 1: Campus Placement Notice (Anonymized Regression Case)
    # Staged compensation, degree percentage, incomplete deadline, circuital branches
    # -------------------------------------------------------------------------
    def test_01_campus_placement_notice_regression(self):
        notice_text = """
        CAMPUS RECRUITMENT NOTICE - 2026 BATCH
        Apex Global Technologies is visiting our campus for hiring Associate Software Engineer.
        
        Job Role: Associate Software Engineer
        Job Type: Full Time (with mandatory 6-month internship)
        
        Eligibility Criteria:
        - B.E / B.Tech in CSE, IT, ECE, EEE
        - Consistent academic record: Minimum 65% in Graduation
        - No active standing backlogs allowed
        
        Compensation Package:
        - Stipend during internship: Rs. 25,000 per month for 6 months
        - CTC upon full-time absorption: 6.5 LPA
        
        Bond / Agreement:
        - Service agreement for 2 years
        
        Work Location: Bangalore / Pune
        
        Required Skills: Java, Spring Boot, SQL, REST APIs, Git
        Preferred Skills: Docker, Angular
        
        Selection Process:
        Round 1: Online Technical Assessment
        Round 2: Technical Interview
        Round 3: HR Interview
        
        Registration Deadline:
        Last date to apply is 15th October at 5:00 PM. No late submissions will be accepted.
        """
        
        result = jd_parser_service.parse_jd_text(notice_text)
        
        # 1. Title & Company
        self.assertIn("Associate Software Engineer", result.title)
        self.assertTrue(result.company_name is not None and "Apex Global Technologies" in result.company_name)
        
        # 2. Academic criteria: 65% in degree, CGPA should NOT be set to 65.0
        self.assertIsNone(result.cgpa_criteria, "65% must not be converted or assigned to cgpa_criteria")
        self.assertIsNotNone(result.academic_percentages)
        self.assertEqual(result.academic_percentages.get("degree"), 65.0)
        
        # 3. Review flag should note academic percentage specified without CGPA
        review_text = " ".join(result.review_flags)
        self.assertIn("percentage", review_text.lower())
        
        # 4. Staged compensation
        self.assertIsNotNone(result.compensation_details)
        self.assertEqual(result.compensation_details.get("stipend"), "₹25,000/month (6 months)")
        self.assertEqual(result.compensation_details.get("full_time_ctc"), "6.5 LPA")
        self.assertIn("Stipend: ₹25,000/month", result.salary_range)
        self.assertIn("Full Time CTC: 6.5 LPA", result.salary_range)
        
        # 5. Incomplete deadline (no year specified)
        self.assertIsNotNone(result.application_deadline)
        self.assertIn("15th October", result.application_deadline)
        # Year must not be hallucinated, and review_flags must flag incomplete date
        self.assertTrue(any("incomplete date" in f.lower() or "year" in f.lower() for f in result.review_flags))
        
        # 6. Branches & Degree
        self.assertTrue(len(result.branch_eligibility) >= 4)
        for b in ["CSE", "IT", "ECE", "EEE"]:
            self.assertIn(b, result.branch_eligibility)
        self.assertIn("B.Tech", result.degree_eligibility or [])
        
        # 7. Bond & Backlogs
        self.assertIsNotNone(result.bond_details)
        self.assertIn("2", result.bond_details)
        self.assertIsNotNone(result.backlog_criteria)
        self.assertIn("no", result.backlog_criteria.lower())
        
        # 8. Skills
        self.assertIn("Java", result.required_skills)
        self.assertIn("Spring Boot", result.required_skills)
        self.assertIn("Docker", result.preferred_skills)

    # -------------------------------------------------------------------------
    # TEST 2: Formal Corporate Enterprise JD
    # Experience, conventional salary range, complete ISO-format deadline
    # -------------------------------------------------------------------------
    def test_02_formal_corporate_jd(self):
        corporate_jd = """
        About CloudScale Solutions:
        CloudScale Solutions is seeking an experienced Senior DevOps Engineer to lead our cloud infrastructure modernization.
        
        Position: Senior DevOps Engineer
        Experience Required: 3-5 years
        Location: Hyderabad (Hybrid)
        Salary Range: ₹14,00,000 - ₹20,00,000 PA
        
        Minimum Eligibility:
        - Bachelor's or Master's degree in Computer Science, Information Technology, or allied branches
        - Minimum CGPA: 7.5 or above
        
        Technical Requirements:
        Must-have skills: Kubernetes, Terraform, AWS, Python, CI/CD, Linux
        Nice-to-have: Helm, Prometheus, Golang
        Certifications: AWS Certified Solutions Architect is a plus.
        
        Application closing date: 2026-11-30 23:59 IST.
        """
        
        result = jd_parser_service.parse_jd_text(corporate_jd)
        
        self.assertEqual(result.title, "Senior DevOps Engineer")
        self.assertEqual(result.experience_level, "3-5 years")
        self.assertEqual(result.work_mode, "Hybrid")
        self.assertIn("Hyderabad", result.location)
        self.assertEqual(result.cgpa_criteria, 7.5)
        self.assertIn("Kubernetes", result.required_skills)
        self.assertIn("Terraform", result.required_skills)
        self.assertIn("AWS", result.required_skills)
        self.assertIn("Golang", result.preferred_skills)
        self.assertIn("2026-11-30", result.application_deadline)
        self.assertIn("14 - 20 LPA", result.salary_range)

    # -------------------------------------------------------------------------
    # TEST 3: Multi-tier Academic Percentages (10th, 12th, Degree)
    # -------------------------------------------------------------------------
    def test_03_multi_tier_academic_percentages(self):
        jd_text = """
        Software Engineer Recruitment at InnovateTech.
        
        Academic Eligibility Criteria:
        - 10th Standard (SSC): Minimum 75%
        - 12th Standard (HSC) or Diploma: Minimum 70%
        - Graduation (B.E/B.Tech): Aggregate of 60% and above
        - Post-Graduation (if applicable): 65% aggregate
        
        Eligible Streams: Computer Science and Engineering
        """
        
        result = jd_parser_service.parse_jd_text(jd_text)
        
        self.assertIsNotNone(result.academic_percentages)
        self.assertEqual(result.academic_percentages.get("10th"), 75.0)
        self.assertEqual(result.academic_percentages.get("12th"), 70.0)
        self.assertEqual(result.academic_percentages.get("degree"), 60.0)
        self.assertEqual(result.academic_percentages.get("postgraduate"), 65.0)
        self.assertIsNone(result.cgpa_criteria)
        self.assertTrue(len(result.field_evidence.get("academic_percentages", [])) >= 2)

    # -------------------------------------------------------------------------
    # TEST 4: CGPA Expression Variations
    # -------------------------------------------------------------------------
    def test_04_cgpa_expressions_variations(self):
        variations = [
            ("Candidates must possess minimum CGPA: 7.0 throughout degree.", 7.0),
            ("Criteria: CGPA >= 8.2 in all semesters with no active arrears.", 8.2),
            ("Cutoff: at least 6.75 CGPA in current engineering course.", 6.75),
            ("Academic performance: Minimum 7.5 CGPA out of 10.", 7.5),
        ]
        
        for text, expected_cgpa in variations:
            with self.subTest(text=text):
                res = jd_parser_service.parse_jd_text(text)
                self.assertEqual(res.cgpa_criteria, expected_cgpa)

    # -------------------------------------------------------------------------
    # TEST 5: Table-based Eligibility Text
    # -------------------------------------------------------------------------
    def test_05_table_based_eligibility(self):
        table_jd = """
        Job Opening: Graduate Trainee
        
        Eligibility Matrix:
        ------------------------------------------
        Level                | Cutoff Required
        ------------------------------------------
        Class X              | 70% and above
        Class XII / PUC      | 65% and above
        Engineering Degree   | 68% or equivalent
        ------------------------------------------
        
        Streams Eligible: CSE, IT, Mechanical, Civil
        """
        
        res = jd_parser_service.parse_jd_text(table_jd)
        self.assertIsNotNone(res.academic_percentages)
        self.assertEqual(res.academic_percentages.get("10th"), 70.0)
        self.assertEqual(res.academic_percentages.get("12th"), 65.0)
        self.assertEqual(res.academic_percentages.get("degree"), 68.0)
        self.assertIn("CSE", res.branch_eligibility)
        self.assertIn("MECH", res.branch_eligibility)

    # -------------------------------------------------------------------------
    # TEST 6: Incomplete vs Complete Dates
    # -------------------------------------------------------------------------
    def test_06_date_handling_incomplete_and_complete(self):
        # Case A: Incomplete date (no year)
        res_incomplete = jd_parser_service.parse_jd_text(
            "Apply by 28th February at 4:00 PM. Applications received after deadline will be rejected."
        )
        self.assertIn("28th February", res_incomplete.application_deadline)
        self.assertTrue(any("year" in f.lower() or "incomplete" in f.lower() for f in res_incomplete.review_flags))
        
        # Case B: Complete date with year
        res_complete = jd_parser_service.parse_jd_text(
            "The last date for student registration is 25th March 2026 at 18:00 hours."
        )
        self.assertIn("2026-03-25", res_complete.application_deadline)
        self.assertFalse(any("missing year" in f.lower() for f in res_complete.review_flags))

    # -------------------------------------------------------------------------
    # TEST 7: Staged Compensation vs Conventional Salary Range
    # -------------------------------------------------------------------------
    def test_07_staged_vs_conventional_compensation(self):
        # Staged compensation
        staged_text = """
        Position: Software Trainee
        During the 3 months internship, an internship stipend of Rs 15,000 per month is offered.
        Post completion of internship, the full-time CTC offered is 5.0 LPA.
        """
        staged_res = jd_parser_service.parse_jd_text(staged_text)
        self.assertIsNotNone(staged_res.compensation_details)
        self.assertIn("staged_compensation", staged_res.compensation_details)
        self.assertEqual(staged_res.compensation_details.get("stipend"), "₹15,000/month (3 months)")
        self.assertEqual(staged_res.compensation_details.get("full_time_ctc"), "5.0 LPA")
        # Ensure it does NOT report a fake range like "15,000 - 5.0"
        self.assertNotIn("15000 - 5", staged_res.salary_range)
        
        # Conventional range
        range_text = """
        Position: Data Scientist
        CTC: 8 - 12 LPA depending on interview performance.
        """
        range_res = jd_parser_service.parse_jd_text(range_text)
        self.assertEqual(range_res.salary_range, "8 - 12 LPA")

    # -------------------------------------------------------------------------
    # TEST 8: Backlog Criteria Variations
    # -------------------------------------------------------------------------
    def test_08_backlog_criteria_parsing(self):
        res1 = jd_parser_service.parse_jd_text("Eligibility: 0 standing arrears / backlogs allowed.")
        self.assertIsNotNone(res1.backlog_criteria)
        self.assertIn("0", res1.backlog_criteria)
        
        res2 = jd_parser_service.parse_jd_text("Students with maximum 1 active backlog may apply.")
        self.assertIsNotNone(res2.backlog_criteria)
        self.assertIn("1", res2.backlog_criteria)

    # -------------------------------------------------------------------------
    # TEST 9: Bond / Service Agreement Extraction
    # -------------------------------------------------------------------------
    def test_09_bond_service_agreement_extraction(self):
        bond_text = """
        Selected candidates have to execute a service agreement of 24 months (2 years)
        with a bond penalty of Rs 1,00,000 if breached.
        """
        res = jd_parser_service.parse_jd_text(bond_text)
        self.assertIsNotNone(res.bond_details)
        self.assertIn("2 years", res.bond_details)

    # -------------------------------------------------------------------------
    # TEST 10: Skills Embedded in Natural Language Paragraphs
    # -------------------------------------------------------------------------
    def test_10_skills_embedded_in_paragraphs(self):
        narrative_jd = """
        We are building high-throughput financial distributed systems. In this role,
        you will be responsible for creating microservices using Go and Python.
        Experience running containerized workloads on Kubernetes with Docker is essential.
        Familiarity with PostgreSQL database querying and Kafka event streaming
        is highly appreciated.
        """
        res = jd_parser_service.parse_jd_text(narrative_jd)
        skills = res.required_skills + res.preferred_skills
        for expected in ["Go", "Python", "Kubernetes", "Docker", "PostgreSQL", "Kafka"]:
            self.assertIn(expected, skills)

    # -------------------------------------------------------------------------
    # TEST 11: Degree Eligibility and Batch Years
    # -------------------------------------------------------------------------
    def test_11_degree_eligibility_and_batch_years(self):
        text = """
        Campus drive for 2025 and 2026 passing out batch.
        Eligible Degrees: B.Tech, M.Tech, and MCA graduates from Computer Science disciplines.
        """
        res = jd_parser_service.parse_jd_text(text)
        self.assertIn("2025", res.graduation_years)
        self.assertIn("2026", res.graduation_years)
        self.assertIn("B.Tech", res.degree_eligibility)
        self.assertIn("M.Tech", res.degree_eligibility)
        self.assertIn("MCA", res.degree_eligibility)

    # -------------------------------------------------------------------------
    # TEST 12: Selection Rounds Extraction
    # -------------------------------------------------------------------------
    def test_12_selection_rounds_extraction(self):
        text = """
        Interview Process:
        Round 1: Cognitive and Technical Assessment
        Round 2: Live System Design & Coding Interview
        Round 3: Behavioral & Culture Fit Round
        """
        res = jd_parser_service.parse_jd_text(text)
        self.assertIsNotNone(res.selection_rounds)
        self.assertTrue(len(res.selection_rounds) >= 3)
        self.assertTrue(any("Assessment" in r for r in res.selection_rounds))
        self.assertTrue(any("Coding" in r or "System Design" in r for r in res.selection_rounds))

    # -------------------------------------------------------------------------
    # TEST 13: Missing and Ambiguous Information Flagging
    # -------------------------------------------------------------------------
    def test_13_missing_and_ambiguous_fields_flagged(self):
        sparse_jd = "Looking for a Frontend Developer with React and JavaScript knowledge."
        res = jd_parser_service.parse_jd_text(sparse_jd)
        
        self.assertIsNone(res.cgpa_criteria)
        self.assertIsNone(res.salary_range)
        self.assertIsNone(res.application_deadline)
        self.assertIsNone(res.academic_percentages)
        self.assertIsNone(res.bond_details)
        
        # Verify evidence tracking exists
        self.assertIn("title", res.field_evidence)
        self.assertIn("required_skills", res.field_evidence)

    # -------------------------------------------------------------------------
    # TEST 14: PDF Document Ingestion (via reportlab)
    # -------------------------------------------------------------------------
    def test_14_pdf_document_ingestion(self):
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=letter)
        c.drawString(100, 750, "CAMPUS HIRING: TechCorp Solutions")
        c.drawString(100, 730, "Role: Software Engineer")
        c.drawString(100, 710, "Minimum CGPA: 7.20")
        c.drawString(100, 690, "Eligible Branches: CSE, IT, ECE")
        c.drawString(100, 670, "Required Skills: Python, Django, SQL")
        c.drawString(100, 650, "Salary: 6 - 9 LPA")
        c.save()
        
        pdf_bytes = buffer.getvalue()
        extracted_text = jd_parser_service.extract_text_from_file(content=pdf_bytes, filename="techcorp_jd.pdf")
        
        self.assertIn("TechCorp Solutions", extracted_text)
        self.assertIn("Software Engineer", extracted_text)
        
        res = jd_parser_service.parse_jd_text(extracted_text)
        self.assertEqual(res.cgpa_criteria, 7.2)
        self.assertIn("Python", res.required_skills)
        self.assertEqual(res.salary_range, "6 - 9 LPA")

    # -------------------------------------------------------------------------
    # TEST 15: DOCX Document Ingestion with Tables (via python-docx)
    # -------------------------------------------------------------------------
    def test_15_docx_document_ingestion_with_tables(self):
        doc = Document()
        doc.add_heading("Quantum Dynamics Hiring Drive", 0)
        doc.add_paragraph("Position: Backend Engineer")
        doc.add_paragraph("Required Skills: Java, Spring Boot, MySQL, Docker")
        
        # Add table
        table = doc.add_table(rows=1, cols=2)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Parameter"
        hdr_cells[1].text = "Value"
        
        row1 = table.add_row().cells
        row1[0].text = "Minimum Degree Aggregate"
        row1[1].text = "68%"
        
        row2 = table.add_row().cells
        row2[0].text = "Eligible Degrees"
        row2[1].text = "B.E, B.Tech, MCA"
        
        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        
        extracted_text = jd_parser_service.extract_text_from_file(content=docx_bytes, filename="quantum_jd.docx")
        self.assertIn("Backend Engineer", extracted_text)
        self.assertIn("68%", extracted_text)
        
        res = jd_parser_service.parse_jd_text(extracted_text)
        self.assertEqual(res.academic_percentages.get("degree"), 68.0)
        self.assertIn("Java", res.required_skills)
        self.assertIn("MCA", res.degree_eligibility)

    # -------------------------------------------------------------------------
    # TEST 16: Plain TXT Document Ingestion
    # -------------------------------------------------------------------------
    def test_16_txt_document_ingestion(self):
        raw_txt = b"Job: AI Research Intern\nSkills: PyTorch, NLP, Python\nLocation: Remote"
        text = jd_parser_service.extract_text_from_file(content=raw_txt, filename="job.txt")
        self.assertIn("AI Research Intern", text)
        
        res = jd_parser_service.parse_jd_text(text)
        self.assertEqual(res.work_mode, "Remote")
        self.assertIn("Python", res.required_skills)

    # -------------------------------------------------------------------------
    # TEST 17: Empty and Malformed Document Error Handling
    # -------------------------------------------------------------------------
    def test_17_empty_and_malformed_document_error_handling(self):
        # Empty file
        with self.assertRaises(ValueError) as ctx:
            jd_parser_service.extract_text_from_file(content=b"", filename="empty.pdf")
        self.assertIn("empty", str(ctx.exception).lower())
        
        # Malformed PDF
        with self.assertRaises(ValueError) as ctx:
            jd_parser_service.extract_text_from_file(content=b"This is not a valid pdf header", filename="corrupt.pdf")
        self.assertIn("unreadable", str(ctx.exception).lower())

    # -------------------------------------------------------------------------
    # TEST 18: End-to-end API Parsing via FastAPI TestClient
    # -------------------------------------------------------------------------
    def test_18_api_parse_jd_text(self):
        from fastapi.testclient import TestClient
        from app.main import app
        from app.api.auth import get_current_placement_officer
        from app.core.database import get_db
        
        po_user = User(id=401, email="po@vcet.ac.in", role=UserRole.PLACEMENT_OFFICER, is_active=True)
        app.dependency_overrides[get_current_placement_officer] = lambda: po_user
        app.dependency_overrides[get_db] = lambda: MagicMock()
        
        try:
            client = TestClient(app)
            payload = {
                "text": "Company: Nexa Corp\nRole: Data Analyst\nRequired Skills: SQL, PowerBI, Excel\nCGPA >= 7.0"
            }
            response = client.post("/placement/jobs/parse-jd", json=payload)
            self.assertEqual(response.status_code, 200)
            data = response.json()
            
            self.assertEqual(data["title"], "Data Analyst")
            self.assertEqual(data["company_name"], "Nexa Corp")
            self.assertEqual(data["cgpa_criteria"], 7.0)
            self.assertIn("SQL", data["required_skills"])
            self.assertIn("cgpa_criteria", data["field_evidence"])
        finally:
            app.dependency_overrides.clear()

    # -------------------------------------------------------------------------
    # TEST 19: End-to-end API File Upload Parsing
    # -------------------------------------------------------------------------
    def test_19_api_upload_jd_file(self):
        from fastapi.testclient import TestClient
        from app.main import app
        from app.api.auth import get_current_placement_officer
        from app.core.database import get_db
        
        po_user = User(id=401, email="po@vcet.ac.in", role=UserRole.PLACEMENT_OFFICER, is_active=True)
        app.dependency_overrides[get_current_placement_officer] = lambda: po_user
        app.dependency_overrides[get_db] = lambda: MagicMock()
        
        try:
            client = TestClient(app)
            doc = Document()
            doc.add_paragraph("Position: Mobile App Developer")
            doc.add_paragraph("Required Skills: Flutter, Dart, Firebase")
            doc.add_paragraph("Minimum CGPA: 6.8")
            buffer = io.BytesIO()
            doc.save(buffer)
            
            files = {
                "file": ("mobile_jd.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            }
            response = client.post("/placement/jobs/upload-jd", files=files)
            self.assertEqual(response.status_code, 200)
            data = response.json()
            
            self.assertEqual(data["title"], "Mobile App Developer")
            self.assertEqual(data["cgpa_criteria"], 6.8)
            self.assertIn("Flutter", data["required_skills"])
        finally:
            app.dependency_overrides.clear()

    # -------------------------------------------------------------------------
    # TEST 20: DRAFT Job Creation Workflow using Parsed Values
    # Placement officer reviews, modifies, and saves as DRAFT
    # -------------------------------------------------------------------------
    def test_20_draft_job_creation_with_parsed_values(self):
        from app.services.placement_service import placement_service
        
        # Simulate extracted values
        parsed = jd_parser_service.parse_jd_text("""
        Recruiter: GlobalTech Systems
        Position: Junior Cloud Engineer
        Eligibility: Minimum 7.0 CGPA, B.E / B.Tech in CSE / IT
        Required Skills: Linux, Python, Bash
        Preferred Skills: AWS, Docker
        Compensation: 6.0 LPA
        """)
        
        # Placement officer prepares create payload from parsed values
        create_payload = JobPostingCreate(
            company_id=1,
            title=parsed.title,
            description="Extracted and verified by Placement Officer",
            job_type="full_time",
            experience_level=parsed.experience_level or "fresher",
            location="Bangalore",
            salary_range=parsed.salary_range,
            required_skills=parsed.required_skills,
            preferred_skills=parsed.preferred_skills,
            cgpa_criteria=parsed.cgpa_criteria,
            branch_eligibility=parsed.branch_eligibility
        )
        
        mock_db = MagicMock()
        created_job = placement_service.create_job_posting(mock_db, create_payload, created_by=1)
        
        # Job must be created in DRAFT status
        self.assertEqual(created_job.status, JobStatus.DRAFT)
        self.assertEqual(created_job.title, "Junior Cloud Engineer")
        self.assertEqual(created_job.cgpa_criteria, 7.0)
        self.assertIn("Linux", created_job.required_skills)


if __name__ == "__main__":
    unittest.main()
