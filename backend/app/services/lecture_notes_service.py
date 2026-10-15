"""
Lecture Notes Ingestion and Lifecycle Service

Handles:
- Streaming upload validation with strict 500 MiB bounded-memory limit
- Magic signature and format validation for PDF, PNG, DOCX, and PPTX
- Deep archive security checks (anti-zip-bomb: entry counts, uncompressed size, compression ratios)
- Resource-bounded multimodal text extraction (pdfplumber with page cap, pytesseract with timeout, python-docx, python-pptx)
- Non-blocking extraction offloaded to a dedicated ThreadPoolExecutor
- Text truncation at 500,000 characters with explicit provenance metadata
- Streaming storage to dedicated 'lecture-notes' MinIO bucket
- Transactional database persistence and failure rollback cleanup
- Atomic conditional claims for retry concurrency
- Note lifecycle updates, publication guards, retries, and safe deletions
"""

import os
import re
import uuid
import zipfile
import logging
import tempfile
import asyncio
from concurrent.futures import ThreadPoolExecutor
from typing import Optional, Tuple, Dict, Any, List
from datetime import datetime, timezone
from PIL import Image
import pytesseract
import pdfplumber
import docx
import pptx

from fastapi import UploadFile, HTTPException, status
from sqlalchemy.orm import Session

from ..core.config import settings
from ..models.user import User, UserRole
from ..models.file import File, FileStatus
from ..models.study_buddy import (
    Subject,
    SyllabusDocument,
    SyllabusUnit,
    SyllabusTopic,
    LectureNote,
    LectureNoteStatus,
    LectureNoteProcessingStatus,
    LectureNoteIndexingStatus,
)
from ..schemas.study_buddy import LectureNoteUpdateRequest
from .file_storage.minio_client import minio_client
from .lecture_notes_indexing_service import lecture_notes_indexing_service

logger = logging.getLogger(__name__)

# ============================================
# EXTRACTION RESOURCE & SAFETY CONSTANTS
# ============================================
MAX_NOTE_FILE_SIZE = 500 * 1024 * 1024  # 500 MiB (524,288,000 bytes)
ALLOWED_NOTE_EXTENSIONS = {".pdf", ".png", ".docx", ".pptx"}
CHUNK_SIZE = 64 * 1024  # 64 KB streaming buffer

# PDF Safeguards
MAX_PDF_PAGES = 300  # Max pages to process per PDF
OCR_TIMEOUT_SECONDS = 10  # Max seconds per pytesseract OCR call
PDF_RENDER_DPI = 150  # Render resolution for OCR fallback

# PNG Safeguards
MAX_PNG_WIDTH = 8192
MAX_PNG_HEIGHT = 8192
MAX_PNG_TOTAL_PIXELS = 25_000_000  # 25 Megapixels

# DOCX & PPTX ZIP Archive Safeguards (Anti-Zip-Bomb)
MAX_ZIP_ENTRIES = 5000  # Max file entries inside the ZIP container
MAX_ZIP_UNCOMPRESSED_SIZE = int(1.5 * 1024 * 1024 * 1024)  # 1.5 GiB uncompressed total
MAX_ZIP_COMPRESSION_RATIO = 100.0  # Max compression ratio for non-empty entries

# Extracted Text Limit
MAX_EXTRACTED_TEXT_LENGTH = 500_000  # 500,000 characters
TRUNCATION_MARKER = (
    "\n\n[TRUNCATED: Content exceeded maximum limit of 500,000 characters. "
    "Remaining content omitted.]"
)

# Dedicated thread pool executor for CPU/IO-bound extraction tasks
# Concurrency limit: 4 workers prevents thread starvation and CPU saturation
extraction_executor = ThreadPoolExecutor(max_workers=4, thread_name_prefix="lecture_note_extractor")


def inspect_zip_safety(temp_path: str, ext: str) -> Optional[str]:
    """
    Inspects ZIP container for ZIP bomb signatures and malformed packages:
    - Entry count <= MAX_ZIP_ENTRIES (5,000)
    - Total uncompressed size <= MAX_ZIP_UNCOMPRESSED_SIZE (1.5 GiB)
    - Per-entry compression ratio <= MAX_ZIP_COMPRESSION_RATIO (100.0x)
    - Expected internal structure ('word/' or 'ppt/')
    Returns an error string if unsafe or malformed, None if safe.
    """
    try:
        with zipfile.ZipFile(temp_path, "r") as z:
            infolist = z.infolist()
            entry_count = len(infolist)
            if entry_count > MAX_ZIP_ENTRIES:
                return f"Archive rejected: entry count ({entry_count}) exceeds safe limit of {MAX_ZIP_ENTRIES}."

            total_uncompressed = 0
            has_structure = False
            expected_prefix = "word/" if ext == ".docx" else "ppt/"

            for info in infolist:
                if info.filename.startswith(expected_prefix):
                    has_structure = True
                total_uncompressed += info.file_size

                # Check compression ratio for entries with compressed size > 0
                if info.compress_size > 0:
                    ratio = info.file_size / info.compress_size
                    if ratio > MAX_ZIP_COMPRESSION_RATIO:
                        return (
                            f"Archive rejected: entry '{info.filename}' has suspicious compression ratio "
                            f"({ratio:.1f}x > {MAX_ZIP_COMPRESSION_RATIO}x threshold)."
                        )

            if not has_structure:
                struct_desc = "Word document" if ext == ".docx" else "PowerPoint presentation"
                return f"Corrupt or invalid {ext.upper()[1:]} file: missing {struct_desc} structure."

            if total_uncompressed > MAX_ZIP_UNCOMPRESSED_SIZE:
                return (
                    f"Archive rejected: total uncompressed size ({total_uncompressed / (1024 * 1024):.1f} MiB) "
                    f"exceeds safe limit of {MAX_ZIP_UNCOMPRESSED_SIZE / (1024 * 1024):.1f} MiB."
                )
    except zipfile.BadZipFile:
        return f"Corrupt or invalid {ext.upper()[1:]} archive: damaged zip package."
    except Exception as e:
        return f"Invalid archive structure: {e.__class__.__name__}."

    return None


class LectureNotesService:
    """Service managing lecture notes ingestion, text extraction, and storage."""

    def __init__(self):
        self.bucket = "lecture-notes"
        self.max_file_size = MAX_NOTE_FILE_SIZE
        self.allowed_extensions = ALLOWED_NOTE_EXTENSIONS
        self.executor = extraction_executor

    def _sanitize_filename(self, filename: Optional[str]) -> str:
        """Sanitize filename to prevent directory traversal and invalid characters."""
        if not filename:
            return "unnamed_lecture_note"
        name = os.path.basename(filename)
        clean = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", name)
        return clean[:200]

    async def _stream_and_validate_upload(
        self,
        file: UploadFile
    ) -> Tuple[str, str, int, str]:
        """
        Streams uploaded file to a temporary file on disk with strict size bounding.
        Validates extension, binary magic signatures, and archive safety.
        Returns: (temp_path, sanitized_filename, total_bytes, mime_type)
        """
        raw_filename = file.filename or "upload"
        sanitized_filename = self._sanitize_filename(raw_filename)
        _, ext = os.path.splitext(sanitized_filename)
        ext = ext.lower()

        if ext not in self.allowed_extensions:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file format '{ext}'. Allowed formats: {', '.join(sorted(self.allowed_extensions))}."
            )

        # Create temporary file
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
        temp_path = temp_file.name
        total_bytes = 0
        header_bytes = b""

        try:
            while True:
                chunk = await file.read(CHUNK_SIZE)
                if not chunk:
                    break
                total_bytes += len(chunk)
                if total_bytes > self.max_file_size:
                    temp_file.close()
                    if os.path.exists(temp_path):
                        os.unlink(temp_path)
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="File exceeds the maximum allowed size of 500 MB."
                    )
                if len(header_bytes) < 32:
                    header_bytes += chunk[: 32 - len(header_bytes)]
                temp_file.write(chunk)

            temp_file.flush()
            temp_file.close()

            if total_bytes == 0:
                if os.path.exists(temp_path):
                    os.unlink(temp_path)
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Empty file cannot be uploaded."
                )

            # Validate binary magic signatures
            if ext == ".pdf":
                if not header_bytes.startswith(b"%PDF-"):
                    if os.path.exists(temp_path):
                        os.unlink(temp_path)
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="Corrupt or invalid PDF file: missing %PDF- magic signature."
                    )
            elif ext == ".png":
                if not header_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
                    if os.path.exists(temp_path):
                        os.unlink(temp_path)
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="Corrupt or invalid PNG file: missing PNG magic signature."
                    )
            elif ext in {".docx", ".pptx"}:
                if not header_bytes.startswith(b"PK\x03\x04"):
                    if os.path.exists(temp_path):
                        os.unlink(temp_path)
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"Corrupt or invalid {ext.upper()[1:]} file: missing PK zip header signature."
                    )

                # Inspect ZIP archive safety (anti-zip-bomb & structural check)
                zip_err = inspect_zip_safety(temp_path, ext)
                if zip_err:
                    if os.path.exists(temp_path):
                        os.unlink(temp_path)
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=zip_err
                    )

            mime_type = file.content_type or "application/octet-stream"
            return temp_path, sanitized_filename, total_bytes, mime_type

        except HTTPException:
            raise
        except Exception as e:
            if os.path.exists(temp_path):
                try:
                    os.unlink(temp_path)
                except Exception:
                    pass
            logger.error(f"Error during streaming upload validation: {e}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Failed to stream and validate uploaded file."
            )

    async def extract_content_async(
        self,
        file_path: str,
        ext: str
    ) -> Tuple[Optional[str], Dict[str, Any], Optional[str]]:
        """
        Asynchronously runs the synchronous, CPU-bound extraction off the main
        asyncio event loop using the dedicated ThreadPoolExecutor.
        """
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(
            self.executor,
            self.extract_content,
            file_path,
            ext
        )

    def extract_content(
        self,
        file_path: str,
        ext: str
    ) -> Tuple[Optional[str], Dict[str, Any], Optional[str]]:
        """
        Extract readable textual content and metadata from the document file on disk.
        Executes with strict safety limits (page limits, image dimension limits,
        OCR timeouts, and 500k text truncation).
        Returns: (extracted_text, metadata, error_msg)
        """
        ext = ext.lower()
        extracted_text = None
        metadata: Dict[str, Any] = {}
        error_msg = None

        try:
            if ext == ".pdf":
                pages_parts = []
                with pdfplumber.open(file_path) as pdf:
                    page_count = len(pdf.pages)
                    metadata["page_count"] = page_count
                    metadata["extractor"] = "pdfplumber"

                    if page_count > MAX_PDF_PAGES:
                        return (
                            None,
                            metadata,
                            f"PDF rejected: page count ({page_count}) exceeds limit of {MAX_PDF_PAGES} pages."
                        )

                    for idx, page in enumerate(pdf.pages):
                        txt = page.extract_text() or ""
                        # OCR fallback on scanned/sparse pages
                        if len(txt.strip()) < 30:
                            try:
                                dpi = PDF_RENDER_DPI
                                if page.width > 2000 or page.height > 2000:
                                    dpi = 72
                                img = page.to_image(resolution=dpi).original
                                ocr_txt = pytesseract.image_to_string(img, timeout=OCR_TIMEOUT_SECONDS)
                                if ocr_txt and len(ocr_txt.strip()) > len(txt.strip()):
                                    txt = ocr_txt
                                    metadata["ocr_pages_used"] = metadata.get("ocr_pages_used", 0) + 1
                            except Exception as ocr_err:
                                logger.debug(f"PDF page {idx + 1} OCR fallback skipped/timed out: {ocr_err}")

                        if txt.strip():
                            pages_parts.append(f"--- Page {idx + 1} ---\n{txt.strip()}")

                full = "\n\n".join(pages_parts).strip()
                if full and len(full) >= 10:
                    if len(full) > MAX_EXTRACTED_TEXT_LENGTH:
                        metadata["truncated"] = True
                        metadata["original_char_count"] = len(full)
                        metadata["truncated_char_count"] = MAX_EXTRACTED_TEXT_LENGTH
                        keep_len = MAX_EXTRACTED_TEXT_LENGTH - len(TRUNCATION_MARKER)
                        extracted_text = full[:keep_len] + TRUNCATION_MARKER
                    else:
                        metadata["truncated"] = False
                        metadata["original_char_count"] = len(full)
                        extracted_text = full
                else:
                    error_msg = "PDF document contained no extractable or readable text content."

            elif ext == ".png":
                metadata["extractor"] = "tesseract_ocr"

                # Step 1: Verify image integrity and decompression bomb check
                try:
                    with Image.open(file_path) as img:
                        img.verify()
                except Image.DecompressionBombError:
                    return None, metadata, "PNG image rejected: exceeds decompression bomb threshold."
                except Exception as e:
                    return None, metadata, f"Corrupt or invalid PNG image: {e.__class__.__name__}."

                # Step 2: Reopen image to inspect dimensions and perform OCR
                try:
                    with Image.open(file_path) as img:
                        w, h = img.width, img.height
                        total_pixels = w * h
                        metadata["dimensions"] = f"{w}x{h}"
                        metadata["total_pixels"] = total_pixels
                        metadata["format"] = img.format

                        if w > MAX_PNG_WIDTH or h > MAX_PNG_HEIGHT:
                            return (
                                None,
                                metadata,
                                f"PNG dimensions ({w}x{h}) exceed maximum allowed {MAX_PNG_WIDTH}x{MAX_PNG_HEIGHT}."
                            )
                        if total_pixels > MAX_PNG_TOTAL_PIXELS:
                            return (
                                None,
                                metadata,
                                f"PNG pixel count ({total_pixels:,}) exceeds maximum allowed {MAX_PNG_TOTAL_PIXELS:,} pixels."
                            )

                        ocr_result = pytesseract.image_to_string(img, timeout=OCR_TIMEOUT_SECONDS)
                except Image.DecompressionBombError:
                    return None, metadata, "PNG image rejected: decompression bomb threshold exceeded."
                except Exception as ocr_err:
                    return None, metadata, f"PNG OCR processing failed or timed out: {ocr_err.__class__.__name__}."

                if ocr_result and len(ocr_result.strip()) >= 10:
                    raw_text = ocr_result.strip()
                    if len(raw_text) > MAX_EXTRACTED_TEXT_LENGTH:
                        metadata["truncated"] = True
                        metadata["original_char_count"] = len(raw_text)
                        metadata["truncated_char_count"] = MAX_EXTRACTED_TEXT_LENGTH
                        keep_len = MAX_EXTRACTED_TEXT_LENGTH - len(TRUNCATION_MARKER)
                        extracted_text = raw_text[:keep_len] + TRUNCATION_MARKER
                    else:
                        metadata["truncated"] = False
                        metadata["original_char_count"] = len(raw_text)
                        extracted_text = raw_text
                else:
                    error_msg = "PNG image contained no discernible text via optical character recognition."

            elif ext == ".docx":
                zip_err = inspect_zip_safety(file_path, ext)
                if zip_err:
                    return None, metadata, zip_err

                doc = docx.Document(file_path)
                metadata["extractor"] = "python-docx"
                metadata["paragraph_count"] = len(doc.paragraphs)
                metadata["table_count"] = len(doc.tables)

                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                table_lines = []
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if row_text:
                            table_lines.append(row_text)

                parts = []
                if paragraphs:
                    parts.extend(paragraphs)
                if table_lines:
                    parts.append("--- Tables ---\n" + "\n".join(table_lines))

                full = "\n\n".join(parts).strip()
                if full and len(full) >= 10:
                    if len(full) > MAX_EXTRACTED_TEXT_LENGTH:
                        metadata["truncated"] = True
                        metadata["original_char_count"] = len(full)
                        metadata["truncated_char_count"] = MAX_EXTRACTED_TEXT_LENGTH
                        keep_len = MAX_EXTRACTED_TEXT_LENGTH - len(TRUNCATION_MARKER)
                        extracted_text = full[:keep_len] + TRUNCATION_MARKER
                    else:
                        metadata["truncated"] = False
                        metadata["original_char_count"] = len(full)
                        extracted_text = full
                else:
                    error_msg = "Word document contained no readable text or paragraph content."

            elif ext == ".pptx":
                zip_err = inspect_zip_safety(file_path, ext)
                if zip_err:
                    return None, metadata, zip_err

                prs = pptx.Presentation(file_path)
                metadata["extractor"] = "python-pptx"
                metadata["slide_count"] = len(prs.slides)

                slide_texts = []
                for s_idx, slide in enumerate(prs.slides):
                    slide_num = s_idx + 1
                    s_parts = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                text_val = paragraph.text.strip()
                                if text_val:
                                    s_parts.append(text_val)
                        elif shape.has_table:
                            for row in shape.table.rows:
                                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                                if row_text:
                                    s_parts.append(row_text)
                    if s_parts:
                        slide_texts.append(f"--- Slide {slide_num} ---\n" + "\n".join(s_parts))

                full = "\n\n".join(slide_texts).strip()
                if full and len(full) >= 10:
                    if len(full) > MAX_EXTRACTED_TEXT_LENGTH:
                        metadata["truncated"] = True
                        metadata["original_char_count"] = len(full)
                        metadata["truncated_char_count"] = MAX_EXTRACTED_TEXT_LENGTH
                        keep_len = MAX_EXTRACTED_TEXT_LENGTH - len(TRUNCATION_MARKER)
                        extracted_text = full[:keep_len] + TRUNCATION_MARKER
                    else:
                        metadata["truncated"] = False
                        metadata["original_char_count"] = len(full)
                        extracted_text = full
                else:
                    error_msg = "PowerPoint presentation contained no readable text frames or slide notes."

        except Exception as e:
            logger.error(f"Text extraction failed for format '{ext}': {e}")
            error_msg = f"Extraction failed: Document is corrupt, encrypted, or invalid ({e.__class__.__name__})."

        return extracted_text, metadata, error_msg

    def validate_unit_topic_alignment(
        self,
        db: Session,
        subject_id: int,
        unit_id: Optional[int],
        topic_id: Optional[int]
    ) -> Tuple[Optional[SyllabusUnit], Optional[SyllabusTopic]]:
        """
        Validates that optional unit_id and topic_id belong to the given subject.
        Raises HTTP 400 on boundary or relationship mismatch.
        """
        unit = None
        topic = None

        if unit_id is not None:
            unit = db.query(SyllabusUnit).join(
                SyllabusDocument, SyllabusUnit.document_id == SyllabusDocument.id
            ).filter(
                SyllabusUnit.id == unit_id,
                SyllabusDocument.subject_id == subject_id
            ).first()

            if not unit:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Unit ID {unit_id} does not belong to subject {subject_id}."
                )

        if topic_id is not None:
            topic_query = db.query(SyllabusTopic).join(
                SyllabusUnit, SyllabusTopic.unit_id == SyllabusUnit.id
            ).join(
                SyllabusDocument, SyllabusUnit.document_id == SyllabusDocument.id
            ).filter(
                SyllabusTopic.id == topic_id,
                SyllabusDocument.subject_id == subject_id
            )

            if unit_id is not None:
                topic_query = topic_query.filter(SyllabusTopic.unit_id == unit_id)

            topic = topic_query.first()
            if not topic:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Topic ID {topic_id} does not belong to the specified unit or subject."
                )

        return unit, topic

    async def upload_and_ingest_note(
        self,
        db: Session,
        subject_id: int,
        file: UploadFile,
        current_user: User,
        title: str,
        description: Optional[str] = None,
        unit_id: Optional[int] = None,
        topic_id: Optional[int] = None
    ) -> LectureNote:
        """
        Stream upload, validate, extract (non-blocking in threadpool), and persist a new lecture note.
        Cleans up temporary disk files and MinIO storage if a failure occurs.
        """
        subject = db.query(Subject).filter(Subject.id == subject_id).first()
        if not subject:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Subject with ID {subject_id} not found."
            )

        if not title or not title.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Title is required and cannot be empty."
            )

        # Validate unit and topic alignment
        self.validate_unit_topic_alignment(db, subject_id, unit_id, topic_id)

        # 1. Stream upload and validate magic bytes and archive safety
        temp_path, sanitized_filename, file_size, mime_type = await self._stream_and_validate_upload(file)
        _, ext = os.path.splitext(sanitized_filename)

        note_uuid = uuid.uuid4()
        object_key = f"notes/{subject.id}/{note_uuid}_{sanitized_filename}"

        minio_uploaded = False
        try:
            # 2. Upload directly from file path to MinIO (streaming multipart)
            minio_client.upload_file_from_path(
                bucket=self.bucket,
                object_key=object_key,
                file_path=temp_path,
                content_type=mime_type
            )
            minio_uploaded = True

            # 3. Create central File audit record
            file_record = File(
                id=note_uuid,
                filename=sanitized_filename,
                file_size=file_size,
                mime_type=mime_type,
                bucket=self.bucket,
                object_key=object_key,
                title=title.strip(),
                description=description or f"Lecture note for {subject.code} - {subject.name}",
                department=subject.department,
                file_type="lecture_note",
                uploaded_by=current_user.id,
                status=FileStatus.ACTIVE
            )
            db.add(file_record)

            # 4. Extract content off the main asyncio event loop via threadpool
            extracted_text, metadata, error_msg = await self.extract_content_async(temp_path, ext)

            if error_msg or not extracted_text:
                processing_status = LectureNoteProcessingStatus.FAILED
                safe_error = error_msg or "Extraction yielded empty or unparseable text."
            else:
                processing_status = LectureNoteProcessingStatus.EXTRACTED
                safe_error = None

            # 5. Persist LectureNote record
            lecture_note = LectureNote(
                subject_id=subject.id,
                uploaded_by_id=current_user.id,
                file_id=file_record.id,
                unit_id=unit_id,
                topic_id=topic_id,
                title=title.strip(),
                description=description.strip() if description else None,
                original_filename=sanitized_filename,
                file_size=file_size,
                mime_type=mime_type,
                minio_bucket=self.bucket,
                minio_object_key=object_key,
                status=LectureNoteStatus.DRAFT,
                is_published=False,
                processing_status=processing_status,
                indexing_status=LectureNoteIndexingStatus.NOT_INDEXED,
                extracted_text=extracted_text,
                extraction_error=safe_error,
                extraction_metadata=metadata,
                indexed_chunk_count=0,
                download_count=0
            )
            db.add(lecture_note)
            db.commit()
            db.refresh(lecture_note)

            logger.info(
                f"✅ Lecture note {lecture_note.id} ({sanitized_filename}) ingested "
                f"for subject {subject.code} (status: {processing_status})"
            )
            return lecture_note

        except HTTPException:
            db.rollback()
            if minio_uploaded:
                try:
                    minio_client.delete_file(self.bucket, object_key)
                except Exception as del_err:
                    logger.warning(f"Failed to cleanup MinIO object on upload abort: {del_err}")
            raise
        except Exception as e:
            db.rollback()
            if minio_uploaded:
                try:
                    minio_client.delete_file(self.bucket, object_key)
                except Exception as del_err:
                    logger.warning(f"Failed to cleanup MinIO object on failure: {del_err}")
            logger.error(f"Error persisting lecture note: {e}")
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Storage or database error occurred while creating lecture note."
            )
        finally:
            if os.path.exists(temp_path):
                try:
                    os.unlink(temp_path)
                except Exception:
                    pass

    async def retry_note_processing(
        self,
        db: Session,
        note_id: int
    ) -> LectureNote:
        """
        Atomically claims an eligible failed note for retry, downloads original file
        from MinIO to a temporary file, and runs non-blocking text extraction.
        Prevents concurrent duplicates and ensures safe state transitions.
        """
        # 1. Atomically claim note: only FAILED or UPLOADED notes are eligible for retry
        rows_updated = db.query(LectureNote).filter(
            LectureNote.id == note_id,
            LectureNote.processing_status.in_([
                LectureNoteProcessingStatus.FAILED,
                LectureNoteProcessingStatus.UPLOADED
            ])
        ).update({
            LectureNote.processing_status: LectureNoteProcessingStatus.PROCESSING,
            LectureNote.updated_at: datetime.now(timezone.utc)
        }, synchronize_session=False)
        db.commit()

        if rows_updated == 0:
            current = db.query(LectureNote).filter(LectureNote.id == note_id).first()
            if not current:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Lecture note with ID {note_id} not found."
                )
            if current.processing_status == LectureNoteProcessingStatus.PROCESSING:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Note is currently being processed."
                )
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Note is not in an eligible state for retry (current status: {current.processing_status.value}). Retry is only allowed for FAILED or UPLOADED notes."
            )

        note = db.query(LectureNote).filter(LectureNote.id == note_id).first()

        _, ext = os.path.splitext(note.original_filename)
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
        temp_path = temp_file.name
        temp_file.close()

        try:
            # 2. Download original file from MinIO directly to temp file path
            minio_client.download_file_to_path(note.minio_bucket, note.minio_object_key, temp_path)

            # 3. Run extraction off the main event loop
            extracted_text, metadata, error_msg = await self.extract_content_async(temp_path, ext)

            # 4. Re-query note in case it was deleted during extraction
            refreshed_note = db.query(LectureNote).filter(LectureNote.id == note_id).first()
            if not refreshed_note:
                logger.warning(f"Note {note_id} was deleted during retry extraction; aborting persistence.")
                return note

            if error_msg or not extracted_text:
                refreshed_note.processing_status = LectureNoteProcessingStatus.FAILED
                refreshed_note.extracted_text = None
                refreshed_note.extraction_error = error_msg or "Extraction yielded empty or unparseable text."
                refreshed_note.extraction_metadata = metadata
            else:
                refreshed_note.processing_status = LectureNoteProcessingStatus.EXTRACTED
                refreshed_note.extracted_text = extracted_text
                refreshed_note.extraction_error = None
                refreshed_note.extraction_metadata = metadata

            db.commit()
            db.refresh(refreshed_note)
            logger.info(f"🔄 Retry processing completed for note {refreshed_note.id} (status: {refreshed_note.processing_status})")
            return refreshed_note

        except Exception as e:
            logger.error(f"Retry processing failed for note {note_id}: {e}")
            db.rollback()
            refreshed_note = db.query(LectureNote).filter(LectureNote.id == note_id).first()
            if refreshed_note:
                refreshed_note.processing_status = LectureNoteProcessingStatus.FAILED
                refreshed_note.extraction_error = f"Retry extraction failed: {e.__class__.__name__}"
                db.commit()
                db.refresh(refreshed_note)
                return refreshed_note
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Retry processing failed: {str(e)}"
            )
        finally:
            if os.path.exists(temp_path):
                try:
                    os.unlink(temp_path)
                except Exception:
                    pass

    def update_note(
        self,
        db: Session,
        note: LectureNote,
        update_data: LectureNoteUpdateRequest
    ) -> LectureNote:
        """
        Updates metadata, optional unit/topic alignment, and publication status.
        Enforces that notes with failed, processing, or missing extractions cannot be published.
        """
        # Validate unit and topic if updated
        new_unit_id = update_data.unit_id if update_data.unit_id is not None else note.unit_id
        new_topic_id = update_data.topic_id if update_data.topic_id is not None else note.topic_id
        if update_data.unit_id is not None or update_data.topic_id is not None:
            self.validate_unit_topic_alignment(db, note.subject_id, new_unit_id, new_topic_id)

        # Publication guard
        is_requesting_publish = (
            update_data.is_published is True or
            update_data.status == LectureNoteStatus.PUBLISHED
        )

        should_trigger_indexing = False
        should_remove_indexing = False

        if is_requesting_publish:
            if note.processing_status != LectureNoteProcessingStatus.EXTRACTED:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot publish lecture note: text extraction must succeed before publication."
                )
            note.is_published = True
            note.status = LectureNoteStatus.PUBLISHED
            note.published_at = datetime.now(timezone.utc)
            should_trigger_indexing = True
        elif update_data.is_published is False:
            note.is_published = False
            if update_data.status:
                note.status = update_data.status
            elif note.status == LectureNoteStatus.PUBLISHED:
                note.status = LectureNoteStatus.DRAFT
            note.indexing_status = LectureNoteIndexingStatus.NOT_INDEXED
            note.indexed_chunk_count = 0
            should_remove_indexing = True

        if update_data.title is not None:
            if not update_data.title.strip():
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Title cannot be empty."
                )
            note.title = update_data.title.strip()

        if update_data.description is not None:
            note.description = update_data.description.strip() if update_data.description else None

        if update_data.unit_id is not None:
            note.unit_id = update_data.unit_id

        if update_data.topic_id is not None:
            note.topic_id = update_data.topic_id

        if update_data.status is not None and not is_requesting_publish:
            note.status = update_data.status
            if update_data.status != LectureNoteStatus.PUBLISHED:
                note.is_published = False
                note.indexing_status = LectureNoteIndexingStatus.NOT_INDEXED
                note.indexed_chunk_count = 0
                should_remove_indexing = True

        db.commit()
        db.refresh(note)

        # Vector indexing lifecycle triggers
        if should_trigger_indexing:
            try:
                lecture_notes_indexing_service.index_lecture_note(db, note.id)
            except Exception as idx_err:
                logger.warning(f"Auto-indexing on publish failed for lecture note {note.id}: {idx_err}")
            db.refresh(note)
        elif should_remove_indexing:
            try:
                lecture_notes_indexing_service.delete_note_chunks(note.id)
            except Exception as del_err:
                logger.warning(f"Failed to delete Qdrant points on unpublish for note {note.id}: {del_err}")

        return note

    def delete_note(
        self,
        db: Session,
        note: LectureNote
    ) -> None:
        """
        Safely removes the lecture note from database and MinIO storage,
        and deletes its vector embeddings from Qdrant.
        """
        note_id = note.id
        bucket = note.minio_bucket
        object_key = note.minio_object_key

        # 1. Clean Qdrant points first
        try:
            lecture_notes_indexing_service.delete_note_chunks(note_id)
        except Exception as qd_err:
            logger.warning(f"Failed to delete Qdrant points for note {note_id}: {qd_err}")

        # 2. Delete database record
        db.delete(note)
        db.commit()

        # 3. Delete object from MinIO
        try:
            minio_client.delete_file(bucket, object_key)
        except Exception as e:
            logger.warning(f"Failed to delete MinIO object {bucket}/{object_key}: {e}")


# Singleton instance
lecture_notes_service = LectureNotesService()
